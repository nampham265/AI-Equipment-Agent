"""Integration tests for role-based authorization and Admin endpoints."""
from __future__ import annotations

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.auth.session import SessionData
from app.core.deps import get_current_user
from app.main import app
from app.repositories.excel_user import ExcelUserRepository
from app.repositories.factory import get_user_repo


@pytest.fixture
def mock_excel_app(tmp_path, monkeypatch):
    """Set up a test environment with a temporary Excel workbook containing Users sheet."""
    wb_path = tmp_path / "auth_test_equipment.xlsx"
    wb = openpyxl.Workbook()
    # Equipment sheet
    ws_eq = wb.active
    ws_eq.title = "Equipment"
    ws_eq.append(["EquipmentID", "EquipmentName", "Category", "Description", "Quantity", "Unit", "Location"])
    ws_eq.append(["EQ-AUTH-01", "Scope", "Tools", "Desc", 5, "unit", "Lab"])

    # Requests sheet
    ws_req = wb.create_sheet("Equipment_Requests")
    ws_req.append([
        "RequestID", "RequesterEmail", "RequesterName", "Project",
        "EquipmentID", "EquipmentName", "Quantity", "NeededFrom", "NeededTo",
        "Purpose", "PriorityLevel", "PriorityReason", "FinalPriorityLevel",
        "PriorityReviewedBy", "PriorityReviewedAt", "Status",
        "ReturnCondition", "ReturnConditionNote", "ReturnedAt",
        "CreatedAt", "UpdatedAt", "RejectionReason"
    ])

    # Users sheet
    ws_u = wb.create_sheet("Users")
    ws_u.append(["UserEmail", "UserName", "Role", "IsActive"])
    ws_u.append(["admin@example.com", "Admin Boss", "ADMIN", "TRUE"])
    ws_u.append(["manager1@example.com", "Manager One", "MANAGER", "TRUE"])
    ws_u.append(["manager2@example.com", "Manager Two", "MANAGER", "TRUE"])
    ws_u.append(["user@example.com", "User Normal", "USER", "TRUE"])
    ws_u.append(["deactivated@example.com", "Banned User", "USER", "FALSE"])

    wb.save(wb_path)
    wb.close()

    from app.core.config import settings
    monkeypatch.setattr(settings, "DATA_BACKEND", "excel")
    monkeypatch.setattr(settings, "EXCEL_DATA_PATH", str(wb_path))

    # Reset repository instances
    import app.repositories.factory as factory
    factory._excel_user_repo = None
    factory._excel_equipment_repo = None
    factory._excel_request_repo = None

    client = TestClient(app)

    def _set_auth(client, email: str, name: str, role: str, is_active: bool = True):
        session_data = SessionData(
            user_email=email,
            user_name=name,
            user_oid=f"oid-{email}",
            role=role,
            is_active=is_active,
        )
        monkeypatch.setattr("app.core.deps.get_session", lambda req: session_data)
        monkeypatch.setattr("app.auth.router.get_session", lambda req: session_data)

    yield client, _set_auth

    app.dependency_overrides.clear()
    factory._excel_user_repo = None
    factory._excel_equipment_repo = None
    factory._excel_request_repo = None


class TestAdminUserManagementEndpoints:
    def test_list_users_as_admin(self, mock_excel_app):
        client, set_auth = mock_excel_app
        set_auth(client, "admin@example.com", "Admin Boss", "ADMIN")
        resp = client.get("/api/admin/users")
        assert resp.status_code == 200
        users = resp.json()
        assert len(users) == 5

    def test_list_users_forbidden_for_user_and_manager(self, mock_excel_app):
        client, set_auth = mock_excel_app
        # User
        set_auth(client, "user@example.com", "User Normal", "USER")
        resp = client.get("/api/admin/users")
        assert resp.status_code == 403

        # Manager
        set_auth(client, "manager1@example.com", "Manager One", "MANAGER")
        resp = client.get("/api/admin/users")
        assert resp.status_code == 403

    def test_admin_create_user(self, mock_excel_app):
        client, set_auth = mock_excel_app
        set_auth(client, "admin@example.com", "Admin Boss", "ADMIN")
        resp = client.post("/api/admin/users", json={
            "user_email": "newbie@example.com",
            "user_name": "New User",
            "role": "USER",
        })
        assert resp.status_code == 201
        data = resp.json()
        assert data["user_email"] == "newbie@example.com"
        assert data["role"] == "USER"

    def test_cannot_create_admin_via_admin_api(self, mock_excel_app):
        client, set_auth = mock_excel_app
        set_auth(client, "admin@example.com", "Admin Boss", "ADMIN")
        resp = client.post("/api/admin/users", json={
            "user_email": "badadmin@example.com",
            "user_name": "Bad Admin",
            "role": "ADMIN",
        })
        # Should be rejected by UserCreate schema validator
        assert resp.status_code == 422

    def test_admin_change_user_role(self, mock_excel_app):
        client, set_auth = mock_excel_app
        set_auth(client, "admin@example.com", "Admin Boss", "ADMIN")
        resp = client.patch("/api/admin/users/user@example.com/role", json={
            "role": "MANAGER",
        })
        assert resp.status_code == 200
        assert resp.json()["role"] == "MANAGER"

    def test_admin_cannot_demote_last_admin(self, mock_excel_app):
        client, set_auth = mock_excel_app
        set_auth(client, "admin@example.com", "Admin Boss", "ADMIN")
        resp = client.patch("/api/admin/users/admin@example.com/role", json={
            "role": "MANAGER",
        })
        assert resp.status_code == 400
        assert "cuối cùng" in resp.json()["detail"]

    def test_admin_toggle_active_status(self, mock_excel_app):
        client, set_auth = mock_excel_app
        set_auth(client, "admin@example.com", "Admin Boss", "ADMIN")
        resp = client.patch("/api/admin/users/user@example.com/active", json={
            "is_active": False,
        })
        assert resp.status_code == 200
        assert resp.json()["is_active"] is False

    def test_admin_cannot_deactivate_last_admin(self, mock_excel_app):
        client, set_auth = mock_excel_app
        set_auth(client, "admin@example.com", "Admin Boss", "ADMIN")
        resp = client.patch("/api/admin/users/admin@example.com/active", json={
            "is_active": False,
        })
        assert resp.status_code == 400
        assert "cuối cùng" in resp.json()["detail"]


class TestManagerSelfApprovalProtection:
    def test_manager_cannot_approve_own_request(self, mock_excel_app):
        client, set_auth = mock_excel_app
        # Manager 1 creates a request
        set_auth(client, "manager1@example.com", "Manager One", "MANAGER")
        create_resp = client.post("/api/requests", json={
            "project": "Manager Project",
            "equipment_id": "EQ-AUTH-01",
            "quantity": 1,
            "needed_from": "2025-08-01",
            "needed_to": "2025-08-05",
            "purpose": "Testing",
        })
        assert create_resp.status_code == 201
        req_id = create_resp.json()["id"]

        # Manager 1 attempts to approve own request -> 403 Forbidden
        approve_resp = client.post(f"/api/manager/requests/{req_id}/approve")
        assert approve_resp.status_code == 403
        assert "chính mình" in approve_resp.json()["detail"]

        # Manager 1 attempts to reject own request -> 403 Forbidden
        reject_resp = client.post(f"/api/manager/requests/{req_id}/reject", json={"reason": "Self"})
        assert reject_resp.status_code == 403

        # Manager 1 attempts to review priority on own request -> 403 Forbidden
        prio_resp = client.patch(f"/api/manager/requests/{req_id}/priority", json={
            "final_priority_level": "High"
        })
        assert prio_resp.status_code == 403

        # Different Manager (Manager 2) CAN approve it -> 200 OK
        set_auth(client, "manager2@example.com", "Manager Two", "MANAGER")
        approve_resp2 = client.post(f"/api/manager/requests/{req_id}/approve")
        assert approve_resp2.status_code == 200
        assert approve_resp2.json()["status"] == "Approved"


class TestDeactivatedUserBlocked:
    def test_deactivated_user_cannot_access(self, mock_excel_app):
        client, set_auth = mock_excel_app
        set_auth(client, "deactivated@example.com", "Banned", "USER", is_active=False)
        resp = client.get("/api/requests/mine")
        assert resp.status_code == 403
        assert "vô hiệu hóa" in resp.json()["detail"]
